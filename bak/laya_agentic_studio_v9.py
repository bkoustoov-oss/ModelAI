import datetime
import os
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox

# Force offline caching
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import customtkinter as ctk
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import laya
import ollama

# Vision Integration
from PIL import Image
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# Configure modern dark theme
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

print("Loading LAYA System 1 Decision Engine (Offline Cache)...")
decision_engine = laya.load("convaiinnovations/laya")

def extract_dual_stream_from_file(file_path, log_callback=None):
    """Extracts spatial and raw streams, with OCR fallback."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ['.png', '.jpg', '.jpeg']:
        text = pytesseract.image_to_string(Image.open(file_path))
        return {"spatial": text, "raw": text}

    elif ext == '.pdf':
        import pdfplumber
        spatial_text = ""
        raw_text = ""
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                s_text = page.extract_text(layout=True)
                r_text = page.extract_text(layout=False)
                
                if not s_text or len(s_text.strip()) < 15:
                    if log_callback: log_callback(f"[{os.path.basename(file_path)}] Page {i+1} is an image. Engaging OCR...")
                    pil_img = page.to_image(resolution=300).original
                    ocr_text = pytesseract.image_to_string(pil_img)
                    s_text = r_text = ocr_text
                
                spatial_text += (s_text or "") + "\n"
                raw_text += (r_text or "") + "\n"
        return {"spatial": spatial_text, "raw": raw_text}
        
    elif ext == '.docx':
        doc = Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return {"spatial": text, "raw": text}
        
    elif ext in ['.xlsx', '.csv']:
        import pandas as pd
        df = pd.read_csv(file_path) if ext == '.csv' else pd.read_excel(file_path)
        text = df.to_string()
        return {"spatial": text, "raw": text}
        
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
            return {"spatial": text, "raw": text}
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")


class AgenticStudioApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Universal LAYA Agentic Studio - v9 (Preview & Responsive UI)")
        self.geometry("1200x900")
        self.minsize(1050, 750)
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=3) # Give right panel more space
        self.grid_rowconfigure(0, weight=1)
        
        self.selected_files = [] 

        # --- LEFT PANEL: CONTROLS (Now Scrollable to fix layout clipping) ---
        self.left_panel = ctk.CTkScrollableFrame(self, corner_radius=10)
        self.left_panel.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        ctk.CTkLabel(self.left_panel, text="Multi-Document Orchestration", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=(10, 15))

        # 1. Multi-File Selection
        ctk.CTkLabel(self.left_panel, text="Source Documents (Click to Preview)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        
        self.file_list_frame = ctk.CTkFrame(self.left_panel, fg_color="#1e1e1e")
        self.file_list_frame.pack(fill="x", padx=10, pady=(5, 5))
        
        btn_frame = ctk.CTkFrame(self.left_panel, fg_color="transparent")
        btn_frame.pack(fill="x", padx=10, pady=(0, 15))
        ctk.CTkButton(btn_frame, text="Add Files...", command=self.add_files, width=120).pack(side="left")
        ctk.CTkButton(btn_frame, text="Clear List", command=self.clear_files, width=120, fg_color="#7f1d1d", hover_color="#991b1b").pack(side="right")

        # 2. Output File Destination
        ctk.CTkLabel(self.left_panel, text="Save Aggregated Output As", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        self.output_path_var = ctk.StringVar()
        self.output_entry = ctk.CTkEntry(self.left_panel, textvariable=self.output_path_var)
        self.output_entry.pack(fill="x", padx=10, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Output", command=self.browse_output).pack(anchor="e", padx=10, pady=(0, 15))
        
        # 3. Temperature Control
        ctk.CTkLabel(self.left_panel, text="AI Temperature (Creativity vs. Strictness)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        self.temp_var = tk.DoubleVar(value=0.2)
        self.temp_slider = ctk.CTkSlider(self.left_panel, from_=0.0, to=1.0, variable=self.temp_var, command=self.update_temp_label)
        self.temp_slider.pack(fill="x", padx=10, pady=(5, 2))
        self.temp_label = ctk.CTkLabel(self.left_panel, text="Current: 0.2 (Strict / Analytical)", font=ctk.CTkFont(size=11, slant="italic"), text_color="#00FF66")
        self.temp_label.pack(anchor="w", padx=10, pady=(0, 15))

        # 4. Directives
        ctk.CTkLabel(self.left_panel, text="Agent Directives", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        self.prompt_text = ctk.CTkTextbox(self.left_panel, height=140)
        self.prompt_text.pack(fill="x", padx=10, pady=(5, 15))
        self.prompt_text.insert("0.0", "Compare the provided documents. Cross-reference the data, highlight discrepancies, and synthesize a unified final report. Preserve original spatial formatting where appropriate.")

        # 5. Action Button
        self.run_btn = ctk.CTkButton(
            self.left_panel, text="▶ Run Multi-File Pipeline", font=ctk.CTkFont(size=14, weight="bold"), 
            height=45, fg_color="#1d4ed8", hover_color="#1e3a8a", command=self.start_pipeline
        )
        self.run_btn.pack(fill="x", padx=10, pady=20)

        # --- RIGHT PANEL: TABBED WORKSPACE ---
        self.right_tabs = ctk.CTkTabview(self, corner_radius=10)
        self.right_tabs.grid(row=0, column=1, padx=(0, 10), pady=10, sticky="nsew")
        
        self.tab_agent = self.right_tabs.add("Agent Workspace")
        self.tab_preview = self.right_tabs.add("Content Preview")

        # Setup Agent Workspace Tab
        self.tab_agent.grid_rowconfigure(0, weight=1)
        self.tab_agent.grid_columnconfigure(0, weight=1)
        self.console = ctk.CTkTextbox(self.tab_agent, fg_color="#0a0a0a", text_color="#d1d5db", font=ctk.CTkFont(family="Consolas", size=12))
        self.console.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        self.console.insert("0.0", "> System ready. Awaiting file batch...\n")

        # Setup Content Preview Tab
        self.tab_preview.grid_rowconfigure(0, weight=1)
        self.tab_preview.grid_columnconfigure(0, weight=1)
        self.preview_box = ctk.CTkTextbox(self.tab_preview, fg_color="#1e1e1e", text_color="#e5e7eb", font=ctk.CTkFont(family="Consolas", size=11))
        self.preview_box.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        self.preview_box.insert("0.0", "Select a file from the left panel to preview its extracted content here.")

    def update_temp_label(self, value):
        val = round(value, 2)
        desc = "Strict" if val <= 0.2 else "Balanced" if val <= 0.5 else "Creative"
        self.temp_label.configure(text=f"Current: {val} ({desc})")

    def add_files(self):
        filenames = filedialog.askopenfilenames(filetypes=[
            ("All Supported Files", "*.docx;*.pdf;*.xlsx;*.csv;*.txt;*.png;*.jpg;*.jpeg")
        ])
        for f in filenames:
            if f not in self.selected_files:
                self.selected_files.append(f)
                # Create a clickable button for each file to trigger the preview
                btn = ctk.CTkButton(
                    self.file_list_frame, 
                    text=f"📄 {os.path.basename(f)}", 
                    anchor="w", 
                    fg_color="transparent", 
                    text_color="#e5e7eb",
                    hover_color="#374151",
                    command=lambda path=f: self.load_preview(path)
                )
                btn.pack(fill="x", pady=2, padx=5)
        
        if self.selected_files and not self.output_path_var.get():
            base_dir = os.path.dirname(self.selected_files[0])
            self.output_path_var.set(os.path.join(base_dir, "Aggregated_Analysis_Report.docx"))

    def load_preview(self, file_path):
        """Extracts and displays the content of the clicked file in the Preview tab."""
        self.right_tabs.set("Content Preview")
        self.preview_box.delete("0.0", "end")
        self.preview_box.insert("end", f"Loading extraction preview for {os.path.basename(file_path)}...\n")
        self.update()
        
        try:
            doc_streams = extract_dual_stream_from_file(file_path)
            
            preview_content = f"=== FILE: {os.path.basename(file_path)} ===\n\n"
            preview_content += "--- 1. SPATIAL LAYOUT STREAM (First 2500 chars) ---\n"
            preview_content += doc_streams['spatial'][:2500]
            if len(doc_streams['spatial']) > 2500: preview_content += "\n...[TRUNCATED FOR PREVIEW]...\n"
            
            preview_content += "\n\n--- 2. SEMANTIC RAW STREAM (First 2500 chars) ---\n"
            preview_content += doc_streams['raw'][:2500]
            if len(doc_streams['raw']) > 2500: preview_content += "\n...[TRUNCATED FOR PREVIEW]...\n"
            
            self.preview_box.delete("0.0", "end")
            self.preview_box.insert("end", preview_content)
        except Exception as e:
            self.preview_box.delete("0.0", "end")
            self.preview_box.insert("end", f"Error extracting file for preview:\n{e}")

    def clear_files(self):
        self.selected_files.clear()
        for widget in self.file_list_frame.winfo_children():
            widget.destroy()
        self.preview_box.delete("0.0", "end")
        self.preview_box.insert("end", "Select a file from the left panel to preview its extracted content here.")

    def browse_output(self):
        filename = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filename:
            self.output_path_var.set(filename)

    def log(self, message, style="system"):
        def _append():
            timestamp = datetime.datetime.now().strftime('%H:%M:%S')
            prefix = f"[{timestamp}] "
            if style == "header":
                msg = f"\n{'-'*50}\n{prefix}⚡ {message}\n{'-'*50}\n"
            elif style == "stream":
                msg = message
            elif style == "success":
                msg = f"\n{prefix}✅ {message}\n"
            else:
                msg = f"\n{prefix} {message}\n"
            
            self.console.insert("end", msg)
            self.console.see("end")
        self.after(0, _append)

    def start_pipeline(self):
        out_path = self.output_path_var.get().strip()
        directives = self.prompt_text.get("0.0", "end").strip()
        temperature = round(self.temp_var.get(), 2)

        if not self.selected_files:
            messagebox.showerror("Error", "Please add at least one input document.")
            return
        if not out_path:
            messagebox.showerror("Error", "Please specify an output save location.")
            return
        
        # Switch to Agent Workspace tab automatically when running
        self.right_tabs.set("Agent Workspace")
        self.run_btn.configure(state="disabled", fg_color="#4b5563")
        threading.Thread(target=self.run_agent, args=(out_path, directives, temperature), daemon=True).start()

    def run_agent(self, out_path, directives, temperature):
        try:
            self.log(f"Extracting and tagging {len(self.selected_files)} files...", "header")
            
            master_spatial = ""
            master_raw = ""
            
            for file_path in self.selected_files:
                file_name = os.path.basename(file_path)
                self.log(f"Processing: {file_name}")
                
                doc_streams = extract_dual_stream_from_file(file_path, log_callback=lambda m: self.log(m, "system"))
                
                master_spatial += f"\n<file name=\"{file_name}\">\n{doc_streams['spatial']}\n</file>\n"
                master_raw += f"\n<file name=\"{file_name}\">\n{doc_streams['raw']}\n</file>\n"

            self.log(f"Agent [GPT OSS 120B] Executing multi-file reasoning... (Temp: {temperature})", "header")
            analysis_prompt = (
                f"Directive: {directives}\n\n"
                "The following data contains multiple files separated by <file> tags. Cross-reference them as requested.\n\n"
                f"=== STREAM 1: SPATIAL LAYOUT ===\n{master_spatial}\n\n"
                f"=== STREAM 2: SEMANTIC RAW TEXT ===\n{master_raw}\n\n"
                "Compare the files based on the directive. Map out how the final document should be structured."
            )
            
            response_stream = ollama.chat(
                model='gpt-oss:120b-cloud',
                messages=[{'role': 'user', 'content': analysis_prompt}],
                options={'num_ctx': 65536, 'temperature': temperature, 'thinking': True},
                stream=True 
            )
            
            mapping_logic = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                mapping_logic += token
                self.log(token, style="stream")

            self.log("\nAgent [LAYA] Routing synthesis strategy...", "header")
            result = decision_engine.predict(
                state=f"GPT Mapping Logic:\n{mapping_logic}",
                questions={
                    "layout_priority": {
                        "type": "choice", 
                        "instructions": "Determine the primary reconstruction strategy based on the multi-file mapping.", 
                        "criteria": {
                            "high_spatial": "Strictly preserve headers and alignments from the source files.",
                            "semantic_repair": "Focus on data aggregation, comparison, and text flow.",
                            "hybrid": "Merge spatial structures with newly synthesized comparative data."
                        }
                    }
                }
            )
            laya_strategy = result.get("answers", result)
            self.log(f"LAYA Router Strategy:\n{laya_strategy}", style="system")

            self.log("\nAgent [Nemotron Ultra] Synthesizing aggregated payload...", "header")
            system_prompt = (
                "You are an advanced multi-document compiler. Output strictly the final document text. "
                "Follow the provided LAYA Strategy to merge the spatial and raw data streams correctly. "
                "CRITICAL: Wrap centered text in <center>...</center>. Wrap right-aligned text in <right>...</right>. "
                "Use **text** for bold. Do not include conversational filler or code blocks."
            )
            
            final_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=[
                    {'role': 'system', 'content': system_prompt}, 
                    {'role': 'user', 'content': f"Rebuild document based on directive ({directives}).\n\nLAYA Strategy: {laya_strategy}\n\nGPT Mapping: {mapping_logic}\n\nRaw Multi-File Data:\n{master_raw}"}
                ],
                options={'num_ctx': 32768, 'temperature': 0.1},
                stream=True
            )
            
            synopsis_text = ""
            for chunk in final_stream:
                token = chunk['message']['content']
                synopsis_text += token
                self.log(token, style="stream")

            self.log("\nConstructing final Word document...", "header")
            out_doc = Document() 
            clean_text = synopsis_text.replace('<br>', '\n').replace('<br/>', '\n')
            
            for line in clean_text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                
                alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                if '<center>' in line or '</center>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.CENTER
                    line = line.replace('<center>', '').replace('</center>', '').strip()
                elif '<right>' in line or '</right>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    line = line.replace('<right>', '').replace('</right>', '').strip()
                
                p = out_doc.add_paragraph()
                p.alignment = alignment
                
                if line.startswith('- ') or line.startswith('* '):
                    line = line[2:].strip()
                    p.add_run("• ").bold = True

                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)

            out_doc.save(out_path)
            self.log(f"Report generated successfully and saved to:\n{out_path}", "success")
            
        except Exception as e:
            self.log(f"PIPELINE ERROR: {str(e)}", "header")
        finally:
            self.after(0, lambda: self.run_btn.configure(state="normal", fg_color="#1d4ed8"))

if __name__ == "__main__":
    app = AgenticStudioApp()
    app.mainloop()