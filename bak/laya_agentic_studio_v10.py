import datetime
import os
import re
import sqlite3
import threading
import uuid
import tkinter as tk
from tkinter import filedialog, messagebox

# Force offline caching
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import customtkinter as ctk
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import laya
import ollama

# Vision Integration
from PIL import Image
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# Configure modern dark theme
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

print("Loading LAYA System 1 Decision Engine (Offline Cache)...")
decision_engine = laya.load("convaiinnovations/laya")

# ==========================================
# MODULE 1: LOCAL SQLITE DATABASE MANAGER
# ==========================================
class DatabaseManager:
    def __init__(self, db_path="agentic_memory.db"):
        self.db_path = db_path
        self._init_db()

    def _init_db(self):
        with sqlite3.connect(self.db_path) as conn:
            conn.execute('''CREATE TABLE IF NOT EXISTS chat_history (
                                id INTEGER PRIMARY KEY AUTOINCREMENT,
                                session_id TEXT,
                                timestamp TEXT,
                                role TEXT,
                                content TEXT
                            )''')
            conn.commit()

    def log_message(self, session_id, role, content):
        timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with sqlite3.connect(self.db_path) as conn:
            conn.execute("INSERT INTO chat_history (session_id, timestamp, role, content) VALUES (?, ?, ?, ?)",
                         (session_id, timestamp, role, content))
            conn.commit()

    def get_session_history(self, session_id):
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.execute("SELECT role, content FROM chat_history WHERE session_id = ? ORDER BY id ASC", (session_id,))
            return [{"role": row[0], "content": row[1]} for row in cursor.fetchall()]

# ==========================================
# MODULE 2: DOCUMENT EXTRACTION ENGINE
# ==========================================
def extract_dual_stream_from_file(file_path, log_callback=None):
    """Extracts spatial and raw streams, with OCR fallback."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ['.png', '.jpg', '.jpeg']:
        text = pytesseract.image_to_string(Image.open(file_path))
        return {"spatial": text, "raw": text}

    elif ext == '.pdf':
        import pdfplumber
        spatial_text, raw_text = "", ""
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                s_text = page.extract_text(layout=True)
                r_text = page.extract_text(layout=False)
                
                if not s_text or len(s_text.strip()) < 15:
                    if log_callback: log_callback(f"[{os.path.basename(file_path)}] Engaging OCR...")
                    pil_img = page.to_image(resolution=300).original
                    s_text = r_text = pytesseract.image_to_string(pil_img)
                
                spatial_text += (s_text or "") + "\n"
                raw_text += (r_text or "") + "\n"
        return {"spatial": spatial_text, "raw": raw_text}
        
    elif ext == '.docx':
        doc = Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return {"spatial": text, "raw": text}
        
    elif ext in ['.xlsx', '.csv']:
        import pandas as pd
        df = pd.read_csv(file_path) if ext == '.csv' else pd.read_excel(file_path)
        return {"spatial": df.to_string(), "raw": df.to_string()}
        
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
            return {"spatial": text, "raw": text}
    else:
        raise ValueError(f"Unsupported file format: {ext}")

# ==========================================
# MODULE 3: WORD COMPILER (Optional Output)
# ==========================================
def compile_word_document(synopsis_text, out_path):
    out_doc = Document() 
    clean_text = synopsis_text.replace('<br>', '\n').replace('<br/>', '\n')
    
    for line in clean_text.split('\n'):
        line = line.strip()
        if not line:
            continue
        
        alignment = WD_ALIGN_PARAGRAPH.LEFT
        if '<center>' in line or '</center>' in line:
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            line = line.replace('<center>', '').replace('</center>', '').strip()
        elif '<right>' in line or '</right>' in line:
            alignment = WD_ALIGN_PARAGRAPH.RIGHT
            line = line.replace('<right>', '').replace('</right>', '').strip()
        
        p = out_doc.add_paragraph()
        p.alignment = alignment
        
        if line.startswith('- ') or line.startswith('* '):
            line = line[2:].strip()
            p.add_run("• ").bold = True

        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                p.add_run(part[2:-2]).bold = True
            else:
                p.add_run(part)

    out_doc.save(out_path)

# ==========================================
# MODULE 4: MAIN APPLICATION GUI
# ==========================================
class AgenticStudioApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Universal LAYA Agentic Studio - v10 (Stateful Chat & DB)")
        self.geometry("1300x900")
        self.minsize(1050, 750)
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=3)
        self.grid_rowconfigure(0, weight=1)
        
        # State Management
        self.db = DatabaseManager()
        self.session_id = str(uuid.uuid4())[:8] # Unique session ID for local history
        self.selected_files = [] 
        self.document_context = "" # Holds the extracted file text for follow-ups
        self.is_processing = False

        self._build_ui()

    def _build_ui(self):
        # --- LEFT PANEL: CONTROLS ---
        self.left_panel = ctk.CTkScrollableFrame(self, corner_radius=10)
        self.left_panel.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        ctk.CTkLabel(self.left_panel, text="Data Sources", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=(10, 15))

        # 1. Multi-File Selection
        self.file_list_frame = ctk.CTkFrame(self.left_panel, fg_color="#1e1e1e")
        self.file_list_frame.pack(fill="x", padx=10, pady=(5, 5))
        
        btn_frame = ctk.CTkFrame(self.left_panel, fg_color="transparent")
        btn_frame.pack(fill="x", padx=10, pady=(0, 15))
        ctk.CTkButton(btn_frame, text="Add Files...", command=self.add_files, width=120).pack(side="left")
        ctk.CTkButton(btn_frame, text="Clear", command=self.clear_files, width=80, fg_color="#7f1d1d", hover_color="#991b1b").pack(side="right")

        # 2. Output File Destination (Now Optional)
        self.export_var = tk.BooleanVar(value=False)
        self.export_chk = ctk.CTkCheckBox(self.left_panel, text="Export to Word Document (.docx)", variable=self.export_var, command=self.toggle_export)
        self.export_chk.pack(anchor="w", padx=10, pady=(15, 5))
        
        self.output_path_var = ctk.StringVar()
        self.output_entry = ctk.CTkEntry(self.left_panel, textvariable=self.output_path_var, state="disabled")
        self.output_entry.pack(fill="x", padx=10, pady=(5, 5))
        self.browse_out_btn = ctk.CTkButton(self.left_panel, text="Browse Output", command=self.browse_output, state="disabled")
        self.browse_out_btn.pack(anchor="e", padx=10, pady=(0, 15))
        
        # 3. Temperature Control
        ctk.CTkLabel(self.left_panel, text="AI Temperature", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        self.temp_var = tk.DoubleVar(value=0.2)
        self.temp_slider = ctk.CTkSlider(self.left_panel, from_=0.0, to=1.0, variable=self.temp_var, command=self.update_temp_label)
        self.temp_slider.pack(fill="x", padx=10, pady=(5, 2))
        self.temp_label = ctk.CTkLabel(self.left_panel, text="Current: 0.2 (Strict / Analytical)", font=ctk.CTkFont(size=11, slant="italic"), text_color="#00FF66")
        self.temp_label.pack(anchor="w", padx=10, pady=(0, 15))

        # 4. Initial Directives
        ctk.CTkLabel(self.left_panel, text="Initial Agent Directives", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=10)
        self.prompt_text = ctk.CTkTextbox(self.left_panel, height=140)
        self.prompt_text.pack(fill="x", padx=10, pady=(5, 15))
        self.prompt_text.insert("0.0", "Compare the provided documents and extract key discrepancies.")

        # 5. Action Button
        self.run_btn = ctk.CTkButton(self.left_panel, text="▶ Initialize Analysis", font=ctk.CTkFont(size=14, weight="bold"), height=45, fg_color="#1d4ed8", hover_color="#1e3a8a", command=self.start_pipeline)
        self.run_btn.pack(fill="x", padx=10, pady=20)

        # --- RIGHT PANEL: INTERACTIVE AGENT WORKSPACE ---
        self.right_tabs = ctk.CTkTabview(self, corner_radius=10)
        self.right_tabs.grid(row=0, column=1, padx=(0, 10), pady=10, sticky="nsew")
        
        self.tab_chat = self.right_tabs.add("Interactive Workspace")
        self.tab_chat.grid_rowconfigure(0, weight=1)
        self.tab_chat.grid_columnconfigure(0, weight=1)

        # Chat History Log
        self.chat_log = ctk.CTkTextbox(self.tab_chat, fg_color="#0a0a0a", text_color="#d1d5db", font=ctk.CTkFont(family="Consolas", size=13))
        self.chat_log.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        self.chat_log.insert("0.0", f"--- Session Initialized [ID: {self.session_id}] ---\nSelect files and click Initialize Analysis to begin.\n\n")

        # Chat Input Area (For follow ups)
        input_frame = ctk.CTkFrame(self.tab_chat, fg_color="transparent")
        input_frame.grid(row=1, column=0, sticky="ew", padx=5, pady=5)
        input_frame.grid_columnconfigure(0, weight=1)

        self.chat_input = ctk.CTkEntry(input_frame, placeholder_text="Ask follow-up questions about the data...", font=ctk.CTkFont(size=13))
        self.chat_input.grid(row=0, column=0, sticky="ew", padx=(0, 10), ipady=8)
        self.chat_input.bind("<Return>", lambda event: self.send_followup())

        self.send_btn = ctk.CTkButton(input_frame, text="Send", font=ctk.CTkFont(weight="bold"), width=80, command=self.send_followup)
        self.send_btn.grid(row=0, column=1, sticky="e", ipady=8)

    # --- UI Helpers ---
    def toggle_export(self):
        state = "normal" if self.export_var.get() else "disabled"
        self.output_entry.configure(state=state)
        self.browse_out_btn.configure(state=state)

    def update_temp_label(self, value):
        val = round(value, 2)
        desc = "Strict" if val <= 0.2 else "Balanced" if val <= 0.5 else "Creative"
        self.temp_label.configure(text=f"Current: {val} ({desc})")

    def add_files(self):
        filenames = filedialog.askopenfilenames(filetypes=[("All Supported Files", "*.docx;*.pdf;*.xlsx;*.csv;*.txt;*.png;*.jpg;*.jpeg")])
        for f in filenames:
            if f not in self.selected_files:
                self.selected_files.append(f)
                ctk.CTkLabel(self.file_list_frame, text=f"📄 {os.path.basename(f)}", font=ctk.CTkFont(size=11), anchor="w").pack(fill="x", pady=2, padx=5)
        if self.selected_files and not self.output_path_var.get():
            self.output_path_var.set(os.path.join(os.path.dirname(self.selected_files[0]), "Agentic_Report.docx"))

    def clear_files(self):
        self.selected_files.clear()
        for widget in self.file_list_frame.winfo_children():
            widget.destroy()

    def browse_output(self):
        filename = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filename: self.output_path_var.set(filename)

    def append_to_chat(self, sender, text, tag=None):
        def _append():
            timestamp = datetime.datetime.now().strftime('%H:%M:%S')
            prefix = f"[{timestamp}] "
            
            if sender == "System":
                msg = f"\n{'-'*40}\n{prefix}⚙️ {text}\n{'-'*40}\n"
            elif sender == "User":
                msg = f"\n{prefix}👤 You:\n{text}\n"
            else:
                msg = f"\n{prefix}🤖 Agent:\n{text}\n" if not tag else text
                
            self.chat_log.insert("end", msg)
            self.chat_log.see("end")
        self.after(0, _append)

    # --- Core Pipeline Execution ---
    def start_pipeline(self):
        if self.is_processing: return
        
        directives = self.prompt_text.get("0.0", "end").strip()
        temperature = round(self.temp_var.get(), 2)

        if not self.selected_files:
            messagebox.showerror("Error", "Please add at least one input document.")
            return
        if self.export_var.get() and not self.output_path_var.get():
            messagebox.showerror("Error", "Please specify an output save location.")
            return
        
        self.is_processing = True
        self.run_btn.configure(state="disabled")
        threading.Thread(target=self.run_initial_agent, args=(directives, temperature), daemon=True).start()

    def run_initial_agent(self, directives, temperature):
        try:
            # 1. EXTRACT DATA & BUILD CONTEXT
            self.append_to_chat("System", f"Extracting and tagging {len(self.selected_files)} files...")
            
            master_raw = ""
            for file_path in self.selected_files:
                doc_streams = extract_dual_stream_from_file(file_path, log_callback=lambda m: self.append_to_chat("System", m))
                master_raw += f"\n<file name=\"{os.path.basename(file_path)}\">\n{doc_streams['raw']}\n</file>\n"
            
            # Store context globally for follow-up chats
            self.document_context = master_raw
            
            # 2. LOG DIRECTIVE TO DB & UI
            self.db.log_message(self.session_id, "user", directives)
            self.append_to_chat("User", directives)

            # 3. LAYA EVALUATION (System 1)
            self.append_to_chat("System", "Consulting LAYA Decision Engine for analytical routing...")
            laya_result = decision_engine.predict(
                state=f"User Request:\n{directives}",
                questions={
                    "strategy": {
                        "type": "choice", 
                        "instructions": "Determine the required analytical approach.", 
                        "criteria": {"data_extraction": "Pulling specific values", "comparison": "Comparing files", "summary": "General overview"}
                    }
                }
            )
            laya_strategy = laya_result.get("answers", laya_result)

            # 4. LLM GENERATION (System 2)
            self.append_to_chat("System", f"Agent synthesizing response (Temp: {temperature})...")
            
            system_prompt = f"You are a data analysis agent. Context files are provided below. Follow the analytical strategy: {laya_strategy}\n\n=== SOURCE DATA ===\n{self.document_context}"
            
            response_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=[{'role': 'system', 'content': system_prompt}, {'role': 'user', 'content': directives}],
                options={'num_ctx': 65536, 'temperature': temperature},
                stream=True
            )
            
            self.append_to_chat("Agent", "") # Print header
            full_response = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                full_response += token
                self.append_to_chat("Agent", token, tag="stream")
                
            self.db.log_message(self.session_id, "agent", full_response)

            # 5. OPTIONAL WORD EXPORT
            if self.export_var.get():
                out_path = self.output_path_var.get().strip()
                self.append_to_chat("System", f"Exporting findings to Word Document...")
                compile_word_document(full_response, out_path)
                self.append_to_chat("System", f"Successfully saved to: {out_path}")

        except Exception as e:
            self.append_to_chat("System", f"ERROR: {str(e)}")
        finally:
            self.is_processing = False
            self.after(0, lambda: self.run_btn.configure(state="normal"))

    # --- Iterative Chat Flow ---
    def send_followup(self):
        if self.is_processing: return
        
        user_text = self.chat_input.get().strip()
        if not user_text: return
        
        if not self.document_context:
            messagebox.showwarning("Warning", "Initialize an analysis with files first before asking follow-ups.")
            return

        self.chat_input.delete(0, "end")
        self.is_processing = True
        threading.Thread(target=self._process_followup, args=(user_text,), daemon=True).start()

    def _process_followup(self, user_text):
        try:
            # 1. UI & DB Updates
            self.append_to_chat("User", user_text)
            self.db.log_message(self.session_id, "user", user_text)
            
            # 2. Rebuild History for LLM Context
            history_records = self.db.get_session_history(self.session_id)
            messages = [{'role': 'system', 'content': f"You are a helpful data agent. Here is the active file context:\n{self.document_context}"}]
            for record in history_records:
                # Map 'agent' role from DB to 'assistant' for Ollama
                role = "assistant" if record["role"] == "agent" else record["role"]
                messages.append({'role': role, 'content': record["content"]})

            # 3. Stream Response
            self.append_to_chat("System", "Thinking...")
            response_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=messages,
                options={'num_ctx': 65536, 'temperature': round(self.temp_var.get(), 2)},
                stream=True
            )
            
            self.append_to_chat("Agent", "")
            full_response = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                full_response += token
                self.append_to_chat("Agent", token, tag="stream")
                
            self.db.log_message(self.session_id, "agent", full_response)

        except Exception as e:
            self.append_to_chat("System", f"ERROR: {str(e)}")
        finally:
            self.is_processing = False

if __name__ == "__main__":
    app = AgenticStudioApp()
    app.mainloop()