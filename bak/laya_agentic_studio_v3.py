import datetime
import os
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox

# Force offline caching
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import customtkinter as ctk
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import laya
import ollama

# Configure modern dark theme
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

print("Loading LAYA System 1 Decision Engine (Offline Cache)...")
decision_engine = laya.load("convaiinnovations/laya")

def add_safe_heading(doc, text, level=1):
    """Adds a stylized heading safely without depending on Word's style dictionary."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(27, 54, 93)   # Deep Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(37, 99, 235)  # Royal Blue
    else:  # Level 3+
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(55, 65, 81)   # Slate Charcoal
    return p

def extract_text_from_file(file_path):
    """Universal parser for different file formats."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.docx':
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    
    elif ext == '.pdf':
        from pypdf import PdfReader
        text = ""
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text
        
    elif ext in ['.xlsx', '.csv']:
        import pandas as pd
        if ext == '.csv':
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)
        # Convert spreadsheet to a string representation
        return df.to_string()
        
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")

class AgenticStudioApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Universal LAYA Agentic Studio - v3")
        self.geometry("1100x780")
        self.minsize(900, 650)
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=2)
        self.grid_rowconfigure(0, weight=1)

        # --- LEFT PANEL: CONTROLS ---
        self.left_panel = ctk.CTkFrame(self, corner_radius=10)
        self.left_panel.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        ctk.CTkLabel(self.left_panel, text="Universal Orchestration", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=(15, 15))

        # 1. Input File Selection
        ctk.CTkLabel(self.left_panel, text="Source Document (Any Format)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.input_path_var = ctk.StringVar()
        self.input_entry = ctk.CTkEntry(self.left_panel, textvariable=self.input_path_var, placeholder_text="Select input file...")
        self.input_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Input", command=self.browse_input).pack(anchor="e", padx=15, pady=(0, 15))

        # 2. Output File Destination
        ctk.CTkLabel(self.left_panel, text="Save Output As (.docx)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.output_path_var = ctk.StringVar()
        self.output_entry = ctk.CTkEntry(self.left_panel, textvariable=self.output_path_var, placeholder_text="Choose save destination...")
        self.output_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Output", command=self.browse_output).pack(anchor="e", padx=15, pady=(0, 15))
        
        # 3. Dynamic Report Title
        ctk.CTkLabel(self.left_panel, text="Output Document Title (Optional)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.title_var = ctk.StringVar(value="Agentic Analysis Report")
        self.title_entry = ctk.CTkEntry(self.left_panel, textvariable=self.title_var, placeholder_text="e.g., Q3 Financial Summary")
        self.title_entry.pack(fill="x", padx=15, pady=(5, 15))

        # 4. Custom Directives Textbox
        ctk.CTkLabel(self.left_panel, text="Agent Directives", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.prompt_text = ctk.CTkTextbox(self.left_panel, height=120)
        self.prompt_text.pack(fill="x", padx=15, pady=(5, 15))
        self.prompt_text.insert("0.0", "Analyze the provided data. Extract key insights and structure them clearly based on the core themes. Do not use markdown tables. Format all data using structured bullet points and bold text.")

        # 5. Action Button
        self.run_btn = ctk.CTkButton(
            self.left_panel, text="▶ Run Agentic Pipeline", font=ctk.CTkFont(size=14, weight="bold"), 
            height=45, fg_color="#1d4ed8", hover_color="#1e3a8a", command=self.start_pipeline
        )
        self.run_btn.pack(fill="x", padx=15, pady=10)

        # --- RIGHT PANEL: AGENT WORKSPACE ---
        self.right_panel = ctk.CTkFrame(self, corner_radius=10, fg_color="#111111")
        self.right_panel.grid(row=0, column=1, padx=(0, 10), pady=10, sticky="nsew")

        ctk.CTkLabel(self.right_panel, text="Agent Workspace (Live Streaming)", text_color="#00FF66", font=ctk.CTkFont(family="Consolas", size=14, weight="bold")).pack(anchor="w", padx=15, pady=10)

        self.console = ctk.CTkTextbox(self.right_panel, fg_color="#0a0a0a", text_color="#d1d5db", font=ctk.CTkFont(family="Consolas", size=12))
        self.console.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.console.insert("0.0", "> System ready. Awaiting multi-format parameters...\n")

    def browse_input(self):
        filename = filedialog.askopenfilename(filetypes=[
            ("All Supported Files", "*.docx;*.pdf;*.xlsx;*.csv;*.txt"),
            ("Word Documents", "*.docx"),
            ("PDF Documents", "*.pdf"),
            ("Excel/CSV Files", "*.xlsx;*.csv"),
            ("Text Files", "*.txt")
        ])
        if filename:
            self.input_path_var.set(filename)
            base, _ = os.path.splitext(filename)
            self.output_path_var.set(f"{base}_Report.docx")

    def browse_output(self):
        filename = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filename:
            self.output_path_var.set(filename)

    def log(self, message, style="system"):
        def _append():
            timestamp = datetime.datetime.now().strftime('%H:%M:%S')
            prefix = f"[{timestamp}] "
            if style == "header":
                msg = f"\n{'-'*50}\n{prefix}⚡ {message}\n{'-'*50}\n"
            elif style == "stream":
                msg = message
            elif style == "success":
                msg = f"\n{prefix}✅ {message}\n"
            else:
                msg = f"\n{prefix} {message}\n"
            
            self.console.insert("end", msg)
            self.console.see("end")
        self.after(0, _append)

    def start_pipeline(self):
        in_path = self.input_path_var.get().strip()
        out_path = self.output_path_var.get().strip()
        directives = self.prompt_text.get("0.0", "end").strip()
        report_title = self.title_var.get().strip()

        if not in_path or not os.path.exists(in_path):
            messagebox.showerror("Error", "Please select a valid input document.")
            return
        if not out_path:
            messagebox.showerror("Error", "Please specify an output save location.")
            return

        self.run_btn.configure(state="disabled", fg_color="#4b5563")
        threading.Thread(target=self.run_agent, args=(in_path, out_path, directives, report_title), daemon=True).start()

    def run_agent(self, in_path, out_path, directives, report_title):
        try:
            # 1. READ DOCUMENT (Universal Parser)
            self.log(f"Extracting data from: {os.path.basename(in_path)}", "header")
            doc_text = extract_text_from_file(in_path)
            
            # Truncate if extreme length to avoid context blowout on massive spreadsheets
            if len(doc_text) > 200000:
                self.log("Document is extremely large. Truncating to fit LLM context limits.")
                doc_text = doc_text[:200000] 
            
            # 2. GENERATE IDEAS (LIVE TOKEN STREAMING)
            self.log("Agent [GPT OSS 120B] Initiating reasoning...", "header")
            response_stream = ollama.chat(
                model='gpt-oss:120b-cloud',
                messages=[{'role': 'user', 'content': f"Directive: {directives}\n\nContext Data:\n{doc_text}"}],
                options={'num_ctx': 65536, 'temperature': 0.7, 'thinking': True},
                stream=True 
            )
            
            raw_ideas = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                raw_ideas += token
                self.log(token, style="stream")

            # 3. LAYA DECISION ENGINE GATE
            self.log("Agent [LAYA] Evaluating insight structure...", "header")
            result = decision_engine.predict(
                state=f"Generated Content:\n{raw_ideas}",
                questions={
                    "is_critical": {"type": "noul", "instructions": "Is this content critical to the overarching directive?"},
                    "category": {"type": "choice", "instructions": "Classify the primary focus.", "criteria": {"data_insight": "Metrics and figures", "strategy": "Planning and execution", "general": "Standard text"}}
                }
            )
            laya_decisions = result.get("answers", result)
            self.log(f"LAYA Router Decisions:\n{laya_decisions}", style="system")

            # 4. FINAL SYNTHESIS
            self.log("Agent [Nemotron Ultra] Synthesizing final report...", "header")
            final_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=[{'role': 'user', 'content': f"Format a polished document based on directive ({directives}):\n\nContent: {raw_ideas}\nLAYA Assessment: {laya_decisions}"}],
                options={'num_ctx': 32768, 'temperature': 0.3},
                stream=True
            )
            
            synopsis_text = ""
            for chunk in final_stream:
                token = chunk['message']['content']
                synopsis_text += token
                self.log(token, style="stream")

            # 5. DOCUMENT WRITING (Generates a brand new .docx report)
            self.log("Constructing final Word document...", "header")
            
            out_doc = Document() # Create a fresh document
            timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            # Apply dynamic title if provided
            if report_title:
                add_safe_heading(out_doc, f"✨ {report_title} ({timestamp})", level=1)

            clean_text = synopsis_text.replace('<br>', '\n').replace('<br/>', '\n')
            
            for line in clean_text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                
                # Headings
                if line.startswith('### '):
                    add_safe_heading(out_doc, line.replace('### ', '').replace('**', ''), level=3)
                elif line.startswith('## '):
                    add_safe_heading(out_doc, line.replace('## ', '').replace('**', ''), level=2)
                elif line.startswith('# '):
                    add_safe_heading(out_doc, line.replace('# ', '').replace('**', ''), level=1)
                
                # Paragraphs & Bullets
                else:
                    p = out_doc.add_paragraph()
                    if line.startswith('- ') or line.startswith('* '):
                        line = line[2:].strip()
                        bullet_run = p.add_run("• ")
                        bullet_run.bold = True
                    
                    parts = re.split(r'(\*\*.*?\*\*)', line)
                    for part in parts:
                        if part.startswith('**') and part.endswith('**'):
                            bold_run = p.add_run(part[2:-2])
                            bold_run.bold = True
                        else:
                            p.add_run(part)

            out_doc.save(out_path)
            self.log(f"Report generated successfully and saved to:\n{out_path}", "success")
            
        except Exception as e:
            self.log(f"PIPELINE ERROR: {str(e)}", "header")
        finally:
            self.after(0, lambda: self.run_btn.configure(state="normal", fg_color="#1d4ed8"))

if __name__ == "__main__":
    app = AgenticStudioApp()
    app.mainloop()