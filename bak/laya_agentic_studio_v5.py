import datetime
import os
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox

# Force offline caching
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import customtkinter as ctk
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH  # NEW: Required for alignment
import laya
import ollama

# Configure modern dark theme
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

print("Loading LAYA System 1 Decision Engine (Offline Cache)...")
decision_engine = laya.load("convaiinnovations/laya")

def extract_text_from_file(file_path):
    """Universal parser for different file formats."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.docx':
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    
    elif ext == '.pdf':
        import pdfplumber
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text(layout=True) 
                if extracted:
                    text += extracted + "\n"
        return text
        
    elif ext in ['.xlsx', '.csv']:
        import pandas as pd
        if ext == '.csv':
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)
        return df.to_string()
        
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")

class AgenticStudioApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Universal LAYA Agentic Studio - v5 (Layout Aware)")
        self.geometry("1100x850")
        self.minsize(950, 700)
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=2)
        self.grid_rowconfigure(0, weight=1)

        # --- LEFT PANEL: CONTROLS ---
        self.left_panel = ctk.CTkFrame(self, corner_radius=10)
        self.left_panel.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        ctk.CTkLabel(self.left_panel, text="Universal Orchestration", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=(15, 15))

        # 1. Input File Selection
        ctk.CTkLabel(self.left_panel, text="Source Document (Any Format)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.input_path_var = ctk.StringVar()
        self.input_entry = ctk.CTkEntry(self.left_panel, textvariable=self.input_path_var, placeholder_text="Select input file...")
        self.input_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Input", command=self.browse_input).pack(anchor="e", padx=15, pady=(0, 15))

        # 2. Output File Destination
        ctk.CTkLabel(self.left_panel, text="Save Output As (.docx)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.output_path_var = ctk.StringVar()
        self.output_entry = ctk.CTkEntry(self.left_panel, textvariable=self.output_path_var, placeholder_text="Choose save destination...")
        self.output_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Output", command=self.browse_output).pack(anchor="e", padx=15, pady=(0, 15))
        
        # 3. Dynamic Report Title
        ctk.CTkLabel(self.left_panel, text="Output Document Title (Optional)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.title_var = ctk.StringVar(value="")
        self.title_entry = ctk.CTkEntry(self.left_panel, textvariable=self.title_var, placeholder_text="Leave blank to skip title injection")
        self.title_entry.pack(fill="x", padx=15, pady=(5, 15))

        # 4. Temperature Control
        ctk.CTkLabel(self.left_panel, text="AI Temperature (Creativity vs. Strictness)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.temp_var = tk.DoubleVar(value=0.1)
        self.temp_slider = ctk.CTkSlider(self.left_panel, from_=0.0, to=1.0, variable=self.temp_var, command=self.update_temp_label)
        self.temp_slider.pack(fill="x", padx=15, pady=(5, 2))
        self.temp_label = ctk.CTkLabel(self.left_panel, text="Current: 0.1 (Strict / Factual)", font=ctk.CTkFont(size=11, slant="italic"), text_color="#00FF66")
        self.temp_label.pack(anchor="w", padx=15, pady=(0, 15))

        # 5. Custom Directives Textbox
        ctk.CTkLabel(self.left_panel, text="Agent Directives", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.prompt_text = ctk.CTkTextbox(self.left_panel, height=120)
        self.prompt_text.pack(fill="x", padx=15, pady=(5, 15))
        self.prompt_text.insert("0.0", "Export the document to a doc format. Keep the original format exactly the same, but execute my instruction: remove the signature.")

        # 6. Action Button
        self.run_btn = ctk.CTkButton(
            self.left_panel, text="▶ Run Agentic Pipeline", font=ctk.CTkFont(size=14, weight="bold"), 
            height=45, fg_color="#1d4ed8", hover_color="#1e3a8a", command=self.start_pipeline
        )
        self.run_btn.pack(fill="x", padx=15, pady=10)

        # --- RIGHT PANEL: AGENT WORKSPACE ---
        self.right_panel = ctk.CTkFrame(self, corner_radius=10, fg_color="#111111")
        self.right_panel.grid(row=0, column=1, padx=(0, 10), pady=10, sticky="nsew")

        ctk.CTkLabel(self.right_panel, text="Agent Workspace (Live Streaming)", text_color="#00FF66", font=ctk.CTkFont(family="Consolas", size=14, weight="bold")).pack(anchor="w", padx=15, pady=10)

        self.console = ctk.CTkTextbox(self.right_panel, fg_color="#0a0a0a", text_color="#d1d5db", font=ctk.CTkFont(family="Consolas", size=12))
        self.console.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.console.insert("0.0", "> System ready. Awaiting multi-format parameters...\n")

    def update_temp_label(self, value):
        val = round(value, 2)
        if val <= 0.2:
            desc = "Strict / Factual"
        elif val <= 0.5:
            desc = "Balanced / Analytical"
        else:
            desc = "Creative / Brainstorming"
        self.temp_label.configure(text=f"Current: {val} ({desc})")

    def browse_input(self):
        filename = filedialog.askopenfilename(filetypes=[("All Supported Files", "*.docx;*.pdf;*.xlsx;*.csv;*.txt")])
        if filename:
            self.input_path_var.set(filename)
            base, _ = os.path.splitext(filename)
            self.output_path_var.set(f"{base}_Report.docx")

    def browse_output(self):
        filename = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filename:
            self.output_path_var.set(filename)

    def log(self, message, style="system"):
        def _append():
            timestamp = datetime.datetime.now().strftime('%H:%M:%S')
            prefix = f"[{timestamp}] "
            if style == "header":
                msg = f"\n{'-'*50}\n{prefix}⚡ {message}\n{'-'*50}\n"
            elif style == "stream":
                msg = message
            elif style == "success":
                msg = f"\n{prefix}✅ {message}\n"
            else:
                msg = f"\n{prefix} {message}\n"
            
            self.console.insert("end", msg)
            self.console.see("end")
        self.after(0, _append)

    def start_pipeline(self):
        in_path = self.input_path_var.get().strip()
        out_path = self.output_path_var.get().strip()
        directives = self.prompt_text.get("0.0", "end").strip()
        report_title = self.title_var.get().strip()
        temperature = round(self.temp_var.get(), 2)

        if not in_path or not os.path.exists(in_path):
            messagebox.showerror("Error", "Please select a valid input document.")
            return
        
        self.run_btn.configure(state="disabled", fg_color="#4b5563")
        threading.Thread(target=self.run_agent, args=(in_path, out_path, directives, report_title, temperature), daemon=True).start()

    def run_agent(self, in_path, out_path, directives, report_title, temperature):
        try:
            # 1. READ DOCUMENT
            self.log(f"Extracting spatial data from: {os.path.basename(in_path)}", "header")
            doc_text = extract_text_from_file(in_path)
            
            # 2. GENERATE IDEAS
            self.log(f"Agent [GPT OSS 120B] Initiating reasoning... (Temp: {temperature})", "header")
            response_stream = ollama.chat(
                model='gpt-oss:120b-cloud',
                messages=[{'role': 'user', 'content': f"Directive: {directives}\n\nContext Data:\n{doc_text}"}],
                options={'num_ctx': 65536, 'temperature': temperature, 'thinking': True},
                stream=True 
            )
            
            raw_ideas = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                raw_ideas += token
                self.log(token, style="stream")

            # 3. FINAL SYNTHESIS WITH LAYOUT PROMPTING
            self.log("\nAgent [Nemotron Ultra] Synthesizing layout-aware report...", "header")
            
            # NEW: Strict System Prompt teaching the AI how to use alignment tags
            system_prompt = (
                "You are an advanced document formatter. Output strictly the final document text. "
                "CRITICAL: You MUST preserve the spatial layout of the original document using tags. "
                "If text was centered, wrap it in <center>...</center>. "
                "If text was on the right, wrap it in <right>...</right>. "
                "Use **text** for bold. Do NOT output markdown code blocks or conversational filler."
            )
            
            final_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=[
                    {'role': 'system', 'content': system_prompt}, 
                    {'role': 'user', 'content': f"Rebuild this document based on directive ({directives}):\n\nContent:\n{raw_ideas}"}
                ],
                options={'num_ctx': 32768, 'temperature': 0.1},
                stream=True
            )
            
            synopsis_text = ""
            for chunk in final_stream:
                token = chunk['message']['content']
                synopsis_text += token
                self.log(token, style="stream")

            # 4. DOCUMENT WRITING & ALIGNMENT PARSING
            self.log("\nConstructing spatially-aware Word document...", "header")
            out_doc = Document() 
            
            if report_title:
                timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                p = out_doc.add_paragraph()
                run = p.add_run(f"✨ {report_title} ({timestamp})")
                run.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = RGBColor(27, 54, 93)
                p.paragraph_format.space_after = Pt(12)

            clean_text = synopsis_text.replace('<br>', '\n').replace('<br/>', '\n')
            
            for line in clean_text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                
                # --- ALIGNMENT PARSING ENGINE ---
                alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Detect and strip <center> tags
                if '<center>' in line or '</center>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.CENTER
                    line = line.replace('<center>', '').replace('</center>', '').strip()
                
                # Detect and strip <right> tags
                elif '<right>' in line or '</right>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    line = line.replace('<right>', '').replace('</right>', '').strip()
                
                # Create paragraph and apply alignment
                p = out_doc.add_paragraph()
                p.alignment = alignment
                
                # Handle bullet points
                if line.startswith('- ') or line.startswith('* '):
                    line = line[2:].strip()
                    p.add_run("• ").bold = True

                # Apply inline bold formatting
                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)

            out_doc.save(out_path)
            self.log(f"Report generated successfully and saved to:\n{out_path}", "success")
            
        except Exception as e:
            self.log(f"PIPELINE ERROR: {str(e)}", "header")
        finally:
            self.after(0, lambda: self.run_btn.configure(state="normal", fg_color="#1d4ed8"))

if __name__ == "__main__":
    app = AgenticStudioApp()
    app.mainloop()