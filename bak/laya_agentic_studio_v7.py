import datetime
import os
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox

# Force offline caching
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import customtkinter as ctk
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import laya
import ollama

# Vision Integration
from PIL import Image
import pytesseract
# Set the Tesseract path for Windows
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# Configure modern dark theme
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

print("Loading LAYA System 1 Decision Engine (Offline Cache)...")
decision_engine = laya.load("convaiinnovations/laya")

def extract_dual_stream_from_file(file_path, log_callback=None):
    """
    Extracts both spatial and raw streams. Includes OCR fallback for scanned documents and images.
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ['.png', '.jpg', '.jpeg']:
        if log_callback: log_callback("Image detected. Initializing OCR Vision Engine...")
        text = pytesseract.image_to_string(Image.open(file_path))
        return {"spatial": text, "raw": text}

    elif ext == '.pdf':
        import pdfplumber
        spatial_text = ""
        raw_text = ""
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                s_text = page.extract_text(layout=True)
                r_text = page.extract_text(layout=False)
                
                # SMART OCR FALLBACK: If standard extraction yields little/no text, it's a scanned page
                if not s_text or len(s_text.strip()) < 15:
                    if log_callback: log_callback(f"Page {i+1} appears to be a scanned image. Engaging OCR...")
                    pil_img = page.to_image(resolution=300).original
                    ocr_text = pytesseract.image_to_string(pil_img)
                    s_text = ocr_text
                    r_text = ocr_text
                
                spatial_text += (s_text or "") + "\n"
                raw_text += (r_text or "") + "\n"
        return {"spatial": spatial_text, "raw": raw_text}
        
    elif ext == '.docx':
        doc = Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return {"spatial": text, "raw": text}
        
    elif ext in ['.xlsx', '.csv']:
        import pandas as pd
        df = pd.read_csv(file_path) if ext == '.csv' else pd.read_excel(file_path)
        text = df.to_string()
        return {"spatial": text, "raw": text}
        
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
            return {"spatial": text, "raw": text}
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")


class AgenticStudioApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Universal LAYA Agentic Studio - v7 (Vision & Layout Aware)")
        self.geometry("1200x850")
        self.minsize(1000, 700)
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=2)
        self.grid_rowconfigure(0, weight=1)

        # --- LEFT PANEL: CONTROLS ---
        self.left_panel = ctk.CTkFrame(self, corner_radius=10)
        self.left_panel.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        ctk.CTkLabel(self.left_panel, text="Universal Orchestration", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=(15, 15))

        # Input/Output
        ctk.CTkLabel(self.left_panel, text="Source Document", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.input_path_var = ctk.StringVar()
        self.input_entry = ctk.CTkEntry(self.left_panel, textvariable=self.input_path_var)
        self.input_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Input", command=self.browse_input).pack(anchor="e", padx=15, pady=(0, 15))

        ctk.CTkLabel(self.left_panel, text="Save Output As", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.output_path_var = ctk.StringVar()
        self.output_entry = ctk.CTkEntry(self.left_panel, textvariable=self.output_path_var)
        self.output_entry.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(self.left_panel, text="Browse Output", command=self.browse_output).pack(anchor="e", padx=15, pady=(0, 15))
        
        # Temperature Control
        ctk.CTkLabel(self.left_panel, text="AI Temperature (Creativity vs. Strictness)", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.temp_var = tk.DoubleVar(value=0.1)
        self.temp_slider = ctk.CTkSlider(self.left_panel, from_=0.0, to=1.0, variable=self.temp_var, command=self.update_temp_label)
        self.temp_slider.pack(fill="x", padx=15, pady=(5, 2))
        self.temp_label = ctk.CTkLabel(self.left_panel, text="Current: 0.1 (Strict / Factual)", font=ctk.CTkFont(size=11, slant="italic"), text_color="#00FF66")
        self.temp_label.pack(anchor="w", padx=15, pady=(0, 15))

        # Directives
        ctk.CTkLabel(self.left_panel, text="Agent Directives", font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=15)
        self.prompt_text = ctk.CTkTextbox(self.left_panel, height=120)
        self.prompt_text.pack(fill="x", padx=15, pady=(5, 15))
        self.prompt_text.insert("0.0", "Rebuild this document as a clean .docx file. Retain the original layout exactly. Synthesize insights if needed.")

        # Action Button
        self.run_btn = ctk.CTkButton(
            self.left_panel, text="▶ Run Agentic Pipeline", font=ctk.CTkFont(size=14, weight="bold"), 
            height=45, fg_color="#1d4ed8", hover_color="#1e3a8a", command=self.start_pipeline
        )
        self.run_btn.pack(fill="x", padx=15, pady=10)

        # --- RIGHT PANEL: AGENT WORKSPACE ---
        self.right_panel = ctk.CTkFrame(self, corner_radius=10, fg_color="#111111")
        self.right_panel.grid(row=0, column=1, padx=(0, 10), pady=10, sticky="nsew")

        ctk.CTkLabel(self.right_panel, text="Agent Workspace (Live Streaming)", text_color="#00FF66", font=ctk.CTkFont(family="Consolas", size=14, weight="bold")).pack(anchor="w", padx=15, pady=10)

        self.console = ctk.CTkTextbox(self.right_panel, fg_color="#0a0a0a", text_color="#d1d5db", font=ctk.CTkFont(family="Consolas", size=12))
        self.console.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.console.insert("0.0", "> System ready. Vision OCR & Dual-Stream parser initialized...\n")

    def update_temp_label(self, value):
        val = round(value, 2)
        desc = "Strict" if val <= 0.2 else "Balanced" if val <= 0.5 else "Creative"
        self.temp_label.configure(text=f"Current: {val} ({desc})")

    def browse_input(self):
        filename = filedialog.askopenfilename(filetypes=[
            ("All Supported Files", "*.docx;*.pdf;*.xlsx;*.csv;*.txt;*.png;*.jpg;*.jpeg"),
            ("Images", "*.png;*.jpg;*.jpeg"),
            ("PDFs", "*.pdf")
        ])
        if filename:
            self.input_path_var.set(filename)
            base, _ = os.path.splitext(filename)
            self.output_path_var.set(f"{base}_Report.docx")

    def browse_output(self):
        filename = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filename:
            self.output_path_var.set(filename)

    def log(self, message, style="system"):
        def _append():
            timestamp = datetime.datetime.now().strftime('%H:%M:%S')
            prefix = f"[{timestamp}] "
            if style == "header":
                msg = f"\n{'-'*50}\n{prefix}⚡ {message}\n{'-'*50}\n"
            elif style == "stream":
                msg = message
            elif style == "success":
                msg = f"\n{prefix}✅ {message}\n"
            else:
                msg = f"\n{prefix} {message}\n"
            
            self.console.insert("end", msg)
            self.console.see("end")
        self.after(0, _append)

    def start_pipeline(self):
        in_path = self.input_path_var.get().strip()
        out_path = self.output_path_var.get().strip()
        directives = self.prompt_text.get("0.0", "end").strip()
        temperature = round(self.temp_var.get(), 2)

        if not in_path or not os.path.exists(in_path):
            messagebox.showerror("Error", "Please select a valid input.")
            return
        
        self.run_btn.configure(state="disabled", fg_color="#4b5563")
        threading.Thread(target=self.run_agent, args=(in_path, out_path, directives, temperature), daemon=True).start()

    def run_agent(self, in_path, out_path, directives, temperature):
        try:
            # 1. DUAL-STREAM INGESTION & VISION OCR
            self.log(f"Extracting Data from: {os.path.basename(in_path)}", "header")
            
            # Pass a callback so the extractor can log OCR activations to the UI
            def ocr_logger(msg):
                self.log(msg, "system")
                
            doc_streams = extract_dual_stream_from_file(in_path, log_callback=ocr_logger)
            
            # 2. GPT OSS ANALYSIS 
            self.log(f"Agent [GPT OSS 120B] Analyzing Spatial vs. Semantic structures...", "header")
            analysis_prompt = (
                f"Directive: {directives}\n\n"
                f"=== STREAM 1: SPATIAL LAYOUT ===\n{doc_streams['spatial']}\n\n"
                f"=== STREAM 2: SEMANTIC RAW TEXT ===\n{doc_streams['raw']}\n\n"
                "Compare these streams. Map out which headers/dates need spatial alignment based on Stream 1, "
                "and identify where paragraph flow needs to be restored based on Stream 2."
            )
            
            response_stream = ollama.chat(
                model='gpt-oss:120b-cloud',
                messages=[{'role': 'user', 'content': analysis_prompt}],
                options={'num_ctx': 65536, 'temperature': temperature, 'thinking': True},
                stream=True 
            )
            
            mapping_logic = ""
            for chunk in response_stream:
                token = chunk['message']['content']
                mapping_logic += token
                self.log(token, style="stream")

            # 3. LAYA DECISION ARBITER
            self.log("\nAgent [LAYA] Routing synthesis strategy...", "header")
            result = decision_engine.predict(
                state=f"GPT Mapping Logic:\n{mapping_logic}",
                questions={
                    "layout_priority": {
                        "type": "choice", 
                        "instructions": "Determine the primary reconstruction strategy based on the mapping.", 
                        "criteria": {
                            "high_spatial": "Strictly preserve headers and center alignments.",
                            "semantic_repair": "Focus heavily on fixing broken sentences and text flow.",
                            "hybrid": "Use spatial for headers, raw for body paragraphs."
                        }
                    }
                }
            )
            laya_strategy = result.get("answers", result)
            self.log(f"LAYA Router Strategy:\n{laya_strategy}", style="system")

            # 4. NEMOTRON FINAL SYNTHESIS
            self.log("\nAgent [Nemotron Ultra] Applying LAYA strategy & generating payload...", "header")
            system_prompt = (
                "You are an advanced document compiler. Output strictly the final document text. "
                "Follow the provided LAYA Strategy to merge the spatial and raw data streams correctly. "
                "CRITICAL: Wrap centered text in <center>...</center>. Wrap right-aligned text in <right>...</right>. "
                "Use **text** for bold. Rebuild broken paragraphs into clean, contiguous sentences. Do not include conversational filler."
            )
            
            final_stream = ollama.chat(
                model='nemotron-3-ultra:cloud',
                messages=[
                    {'role': 'system', 'content': system_prompt}, 
                    {'role': 'user', 'content': f"Rebuild document based on directive ({directives}).\n\nLAYA Strategy: {laya_strategy}\n\nGPT Mapping: {mapping_logic}\n\nRaw Data:\n{doc_streams['raw']}"}
                ],
                options={'num_ctx': 32768, 'temperature': 0.1},
                stream=True
            )
            
            synopsis_text = ""
            for chunk in final_stream:
                token = chunk['message']['content']
                synopsis_text += token
                self.log(token, style="stream")

            # 5. DOCUMENT WRITER
            self.log("\nConstructing final Word document...", "header")
            out_doc = Document() 
            clean_text = synopsis_text.replace('<br>', '\n').replace('<br/>', '\n')
            
            for line in clean_text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                
                alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                if '<center>' in line or '</center>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.CENTER
                    line = line.replace('<center>', '').replace('</center>', '').strip()
                elif '<right>' in line or '</right>' in line:
                    alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    line = line.replace('<right>', '').replace('</right>', '').strip()
                
                p = out_doc.add_paragraph()
                p.alignment = alignment
                
                if line.startswith('- ') or line.startswith('* '):
                    line = line[2:].strip()
                    p.add_run("• ").bold = True

                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)

            out_doc.save(out_path)
            self.log(f"Report generated successfully and saved to:\n{out_path}", "success")
            
        except Exception as e:
            self.log(f"PIPELINE ERROR: {str(e)}", "header")
        finally:
            self.after(0, lambda: self.run_btn.configure(state="normal", fg_color="#1d4ed8"))

if __name__ == "__main__":
    app = AgenticStudioApp()
    app.mainloop()